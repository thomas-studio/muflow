"""Export canonical assets through the user's v2 sheet contract, with exact-ID traces."""
from __future__ import annotations
import argparse, collections, copy, json
from pathlib import Path
from core import load,save
from asset_common import sheet_plan,assemble_extra_prompt,assemble_group_prompt
from asset_workbook import build,fill_prompt

def export(data,project,out,refresh=False):
    data=copy.deepcopy(data);out=Path(out);out.mkdir(parents=True,exist_ok=True)
    targets=['资产拆解表.xlsx','资产清单源.json','生图提示词清单.md','资产图_出图任务.json']
    if any((out/n).exists() for n in targets):
        if not refresh:raise ValueError('Use a new export version directory or explicit local --refresh')
        import datetime,shutil
        archive=out/'history'/datetime.datetime.now(datetime.timezone.utc).strftime('%Y%m%dT%H%M%S%f')
        archive.mkdir(parents=True)
        for name in targets:
            if (out/name).is_file():shutil.copyfile(out/name,archive/name)
    cfg=project.get('settings',{});medium=cfg.get('medium');quality=cfg.get('quality');ratio=cfg.get('scene_ratio')
    if not medium or not quality or ratio not in ('16:9','21:9'):raise ValueError('Confirmed medium/quality/scene_ratio required')
    assets={a['id']:a for a in project.get('assets',[])}
    for row in data.get('表6 一次性人物清单',{}).get('rows',[]):
        if not row.get('简描(给AI出图参考)') and row.get('简描'):row['简描(给AI出图参考)']=row['简描']
        if not row.get('简描(给AI出图参考)') and '不调用人物图' not in row.get('备注',''):raise ValueError('Missing extra description: '+row.get('编号',''))
    usages=collections.defaultdict(set)
    for g in project.get('groups',[]):
        for shot in g.get('shots',[]):
            for aid in set(shot.get('asset_ids',[])):
                if aid not in assets:raise ValueError('Unknown asset '+aid)
                usages[aid].add(shot['scene_id'])
    ordered=project.get('scenes',[]);scene_map={s['id']:s for s in ordered};trace={}
    for aid,ss in usages.items():
        scenes=[s for s in ordered if s['id'] in ss];eps=list(dict.fromkeys(int(s['episode']) for s in scenes))
        trace[aid]={'首次出现':('第%02d集%s场'%(scenes[0]['episode'],scenes[0]['id'])) if scenes else '',
                    '后续复用集数':'、'.join('%02d'%x for x in eps[1:]),
                    '总复用次数':f'画面调用共{len(scenes)}场（{len(eps)}集）'}
    t3=[]
    for sc in ordered:
        visible=[aid for aid in assets if sc['id'] in usages[aid]]
        if not visible:continue
        action=sc.get('key_action','')
        if not 20<=len(action)<=120:raise ValueError('key_action must be 20–120 visible-action characters: '+sc['id'])
        scene_ids=[a for a in visible if assets[a]['kind']=='scene'];roles=[a for a in visible if assets[a]['kind'] in ('role','extra','group')]
        t3.append({'集':'第%02d集'%sc['episode'],'场':sc['id'],'场景':sc['name'],
                   '场景标注':'；'.join(scene_ids),'出场角色':'；'.join(assets[a]['name'] for a in roles),
                   '角色标注':'；'.join(roles),'关键动作':action})
    data['表3 分集索引表']={'口径':'由实际镜头与完整资产版本ID反算，声音角色不计人物图调用','rows':t3}
    for label,kind,namekey in [('表4 资产溯源-场景','scene','场景名'),('表5 资产溯源-角色','role','角色名')]:
        rows=[]
        for aid,a in assets.items():
            if a['kind']!=kind or aid not in trace:continue
            t=trace[aid];r={'资产编号':aid,namekey:a['name'],'类型':'制作资产',**t,
              '备注':('绑定母图 '+a['parent_id']) if a.get('parent_id') else '独立基线'}
            if kind=='scene':r['需要的变体版本']=aid
            else:r['需要的服化变体']=aid;r['总复用次数']=t['总复用次数'].replace('画面调用','人物图调用')
            rows.append(r)
        data[label]={'口径':'精确完整ID统计，每个阶段独立首次出现','rows':rows}
    tasks=[];prompts=[];seen=set()
    specs=[('表1 核心角色卡','姓名'),('表2 核心场景卡','场景'),('表7 道具与特殊资产','资产名'),('表6 一次性人物清单','人物')]
    for sheet,namecol in specs:
        for row in data.get(sheet,{}).get('rows',[]):
            aid=row.get('编号');a=assets.get(aid)
            if not a:raise ValueError('Source row not in manifest assets: '+str(aid))
            if aid in seen:raise ValueError('Duplicate asset row '+aid)
            seen.add(aid)
            if a.get('voice_only'):continue
            if aid in trace:row.update({k:v for k,v in trace[aid].items() if k in row or sheet=='表7 道具与特殊资产'})
            p=row.get('生图提示词') or a.get('prompt')
            if not p:
                if a.get('mode')=='edit' or a.get('template')=='custom':raise ValueError('Explicit edit/custom prompt required: '+aid)
                if sheet=='表6 一次性人物清单':
                    group=a['kind']=='group'
                    p=(assemble_group_prompt if group else assemble_extra_prompt)(row[namecol],row,medium,quality)
                else:p=fill_prompt(sheet,row,medium,quality,ratio)
            if not p:raise ValueError('Missing prompt '+aid)
            if sheet!='表6 一次性人物清单':row['生图提示词']=p
            mode=a.get('mode','generate');parent=a.get('parent_id')
            if parent and mode!='edit':raise ValueError('Derived assets must use actual image editing: '+aid)
            if mode=='edit' and parent not in assets:raise ValueError('Edit mother missing '+aid)
            tasks.append({'id':aid,'name':a['name'],'operation':'reuse' if mode=='reuse' else 'image_edit' if mode=='edit' else 'image_generate','generation_required':mode!='reuse','local_path':a.get('local_path'),'sha256':a.get('sha256'),
                'parent_id':parent,'depends_on':[parent] if mode=='edit' else [],'prompt':p,
                'first_seen':trace.get(aid,{}).get('首次出现',''),'ratio':ratio if a['kind']=='scene' else '16:9'})
            prompts.append(f'{aid}｜{a["name"]}｜{mode}\n{p}\n')
    data.update({'媒介':medium,'统一画质串':quality,'场景画幅':ratio})
    warnings=build(data,str(out/targets[0]),medium,quality,ratio)
    save(out/targets[1],data);(out/targets[2]).write_text('\n'.join(prompts),encoding='utf-8')
    save(out/targets[3],{'items':tasks,'warnings_count':warnings,'billing':'Actual configured provider quotation; no assumed unit price'})
    # Read back every written source field; template-generated prompt cells are in data above.
    from openpyxl import load_workbook
    wb=load_workbook(out/targets[0],data_only=True);cells=0
    for label,headers in sheet_plan(bool(data.get('表7 道具与特殊资产'))):
        ws=wb[label]
        for ri,row in enumerate(data.get(label,{}).get('rows',[]),4):
            for ci,key in enumerate(headers,1):
                if str(ws.cell(ri,ci).value or '')!=str(row.get(key) or ''):raise ValueError('Workbook mismatch: '+label+' '+key)
                cells+=1
    return {'assets':len(tasks),'edits':sum(t['operation']=='image_edit' for t in tasks),'verified_cells':cells,'warnings':warnings}
if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('source');p.add_argument('project');p.add_argument('out');p.add_argument('--refresh',action='store_true');a=p.parse_args()
    print(json.dumps(export(load(a.source),load(a.project),a.out,a.refresh),ensure_ascii=False))
