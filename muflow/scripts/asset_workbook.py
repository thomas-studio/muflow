#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""从清单 JSON 生成合规的资产拆解表（含「生图提示词」列）。

用法：
    python build_assets_xlsx.py <清单.json> <输出.xlsx>
        [--medium 真人] [--quality "统一画质串"] [--scene-ratio 16:9|21:9]

清单 JSON 结构（Sheet 名用固定名；rows 是「列名 -> 值」的字典，列名照抄固定表头）：
{
  "剧名": "示例作品",
  "媒介": "真人",
  "统一画质串": "成熟写实人像摄影质感，皮肤纹理与发丝细节真实，中焦段自然透视，色彩还原自然",
  "场景画幅": "16:9",
  "表1 核心角色卡": {"口径": "…", "rows": [{"编号":"A1","姓名":"示例角色", "...":"..."}]},
  "表2 核心场景卡": {"口径": "…", "rows": [ … ]},
  "表3 分集索引表": {"口径": "…", "rows": [ … ]},
  "表4 资产溯源-场景": {"口径": "…", "rows": [ … ]},
  "表5 资产溯源-角色": {"口径": "…", "rows": [ … ]},
  "表6 一次性人物清单": {"口径": "…", "rows": [ … ]},
  "表7 道具与特殊资产": {"口径": "…", "rows": [ … ]},   ← 可选，给了才生成这张表
  "出图风险提醒": {"口径": "…", "rows": [ … ]}
}

表1/表2/表7 的「生图提示词」若留空，脚本按结构化字段自动拼装成稿句。
"""
from __future__ import annotations

import argparse
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from asset_common import (  # noqa: E402
    DATA_START, DEFAULT_MEDIUM, DEFAULT_QUALITY, HEADER_ROW, OPTIONAL_SHEET_PROPS,
    RATIO_SCENE_DEFAULT, RATIO_SCENE_OPTIONS, SCENE_RATIO_NOTICE,
    angle_leak, assemble_extra_prompt, assemble_group_prompt, assemble_prop_prompt,
    assemble_role_edit_prompt, assemble_role_prompt, assemble_scene_prompt,
    assemble_special_prompt, banned_layout_hits, is_stage_version, sheet_plan,
)

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font  # noqa: E402


def fill_prompt(sheet_name: str, row: dict, medium: str, quality: str, scene_ratio: str) -> str:
    d = {k: ("" if v is None else str(v)) for k, v in row.items()}
    if sheet_name == "表1 核心角色卡":
        if is_stage_version(d):
            return assemble_role_edit_prompt(d.get("姓名", ""), d, medium, quality)
        return assemble_role_prompt(d.get("姓名", ""), d, medium, quality)
    if sheet_name == "表2 核心场景卡":
        return assemble_scene_prompt(d.get("场景", ""), d, medium, quality, scene_ratio)
    if sheet_name == OPTIONAL_SHEET_PROPS[0]:
        t = (d.get("类型") or "").strip()
        nm = d.get("资产名", "")
        if t in ("道具", "道具类"):
            return assemble_prop_prompt(nm, d, medium, quality)
        return assemble_special_prompt(nm, d, medium, quality)
    return ""


def build(data: dict, out_path: str, medium: str, quality: str, scene_ratio: str) -> int:
    title_prefix = data.get("剧名") or "未命名"
    with_props = bool(data.get(OPTIONAL_SHEET_PROPS[0]))
    plan = sheet_plan(with_props)

    wb = Workbook()
    wb.remove(wb.active)
    warns: list[str] = []

    for sheet_name, headers in plan:
        block = data.get(sheet_name) or {}
        rows = block.get("rows") or []
        note = block.get("口径") or ""

        ws = wb.create_sheet(sheet_name)
        ws.cell(row=1, column=1, value=f"{title_prefix}｜{sheet_name}").font = Font(bold=True, size=12)
        ws.cell(row=2, column=1, value=note)

        for c, h in enumerate(headers, 1):
            cell = ws.cell(row=HEADER_ROW, column=c, value=h)
            cell.font = Font(bold=True)

        for i, row in enumerate(rows):
            r = DATA_START + i
            for c, h in enumerate(headers, 1):
                v = row.get(h, "")
                if h == "生图提示词" and not str(v or "").strip():
                    v = fill_prompt(sheet_name, row, medium, quality, scene_ratio)
                ws.cell(row=r, column=c, value=v)
            # 提示词自检
            p = str(ws.cell(row=r, column=headers.index("生图提示词") + 1).value or "") \
                if "生图提示词" in headers else ""
            if p:
                bad = banned_layout_hits(p)
                if bad:
                    warns.append(f"{sheet_name} 行{r} 提示词出现栏位结构词：{'、'.join(bad)}（会被画成分隔线）")
                nm = row.get("姓名") or row.get("场景") or row.get("资产名") or ""
                if sheet_name == "表1 核心角色卡" and "肖像" in p:
                    tail = p.split("。", 2)[-1]
                    leak = angle_leak(tail)
                    if leak:
                        warns.append(f"{sheet_name} 行{r}「{nm}」固定顺序之后的正文里残留角度词："
                                     f"{'、'.join(leak)}（可能被当成新增人像）")

        last_row = max(DATA_START - 1, DATA_START + len(rows) - 1)
        last_col = len(headers)
        ws.freeze_panes = "A4"
        ws.auto_filter.ref = f"A{HEADER_ROW}:{_col_letter(last_col)}{last_row}"
        for r in range(DATA_START, last_row + 1):
            ws.row_dimensions[r].height = None
            for c in range(1, last_col + 1):
                ws.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")
        for c in range(1, last_col + 1):
            ws.column_dimensions[_col_letter(c)].width = 22
        ws.column_dimensions["A"].width = 10

    os.makedirs(os.path.dirname(os.path.abspath(out_path)) or ".", exist_ok=True)
    wb.save(out_path)
    print(f"[OK] 已生成 {out_path}")
    print(f"     全剧统一风格：{medium}画面，{quality}")
    print(f"     画幅：角色/道具/特殊资产/群像 = 16:9（固定）；场景 = {scene_ratio}")
    for sheet_name, headers in plan:
        n = len((data.get(sheet_name) or {}).get("rows") or [])
        print(f"     {sheet_name}：{len(headers)}列 / {n}行")
    if warns:
        print("\n⚠ 提示词自检发现问题（必须改完再出图）：")
        for x in warns:
            print("  ·", x)
    return len(warns)


def _col_letter(n: int) -> str:
    s = ""
    while n > 0:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def main() -> int:
    ap = argparse.ArgumentParser(description="从清单 JSON 生成合规资产拆解表")
    ap.add_argument("json_path")
    ap.add_argument("out_path")
    ap.add_argument("--medium", default=None,
                    help="全剧画面媒介（真人/二维动画/三维动画/定格动画/水墨动画/黏土动画）")
    ap.add_argument("--quality", default=None, help="全剧统一画质串（逐字相同）")
    ap.add_argument("--scene-ratio", default=None, choices=list(RATIO_SCENE_OPTIONS),
                    help="场景画幅；**必须先经用户确认**，未给按 16:9 并在交付时提示确认")
    a = ap.parse_args()

    if not os.path.isfile(a.json_path):
        print(f"[ERROR] 清单不存在：{a.json_path}")
        return 2
    data = json.load(open(a.json_path, encoding="utf-8"))
    medium = a.medium or data.get("媒介") or data.get("风格") or DEFAULT_MEDIUM
    quality = a.quality if a.quality is not None else (data.get("统一画质串") or DEFAULT_QUALITY)
    scene_ratio = a.scene_ratio or data.get("场景画幅") or RATIO_SCENE_DEFAULT
    if not (a.scene_ratio or data.get("场景画幅")):
        print(f"[需确认] {SCENE_RATIO_NOTICE}")
    build(data, a.out_path, medium, quality, scene_ratio)
    return 0


if __name__ == "__main__":
    sys.exit(main())
